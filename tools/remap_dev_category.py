# -*- coding: utf-8 -*-
"""개발이력 시트의 '구분'을 자동화 개발 / 보고서 생성 두 가지로 바꾼다.

'주요 프로젝트 / 업무 효율화'는 중요도 구분이라 홈페이지에서 쓰기 애매했다.
무엇을 만들어 주는 회사인지 드러나도록 결과물 종류로 바꾼다.

제자리에서 고치므로 입력한 내용은 그대로 남는다. (최초 1회용)
"""
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

import sys
sys.path.insert(0, str(Path(__file__).parent))
from 엑셀경로 import 엑셀_찾기

엑셀_경로 = 엑셀_찾기()
마지막_행 = 300

# 결과물이 '문서·보고서'인 건들. 나머지는 모두 자동화 개발로 본다.
보고서_건 = {
    "역량 측정 보고서 자동 생성 프로그램 개발",
    "전공역량진단평가 분석 시스템",
    "교육 운영 후 통계 자료 자동 생성",
    "주간 업무 보고 자료 자동 생성",
    "만족도 조사 결과 보고서 내용 자동 작성",
    "전체 수료증 자동 출력 파일",
    "업무 매뉴얼 제작",
}

통합문서 = load_workbook(엑셀_경로)
시트 = 통합문서["개발이력"]
열 = {str(칸.value).strip(): 칸.column for 칸 in 시트[4] if 칸.value}

바뀐_수 = 0
for 행 in range(5, 시트.max_row + 1):
    이름 = 시트.cell(행, 열["프로젝트명"]).value
    if not 이름:
        continue
    비고 = str(시트.cell(행, 열["비고"]).value or "")

    새_구분 = "보고서 생성" if str(이름).strip() in 보고서_건 else "자동화 개발"
    옛_구분 = 시트.cell(행, 열["구분"]).value

    # 예시 행도 형식을 맞춰 둔다
    시트.cell(행, 열["구분"], 새_구분)
    if 옛_구분 != 새_구분:
        바뀐_수 += 1
        표식 = "예시" if 비고.startswith("예시") else f"[{시트.cell(행, 열['홈페이지표시']).value}]"
        print(f"  {행:>3} {표식} {옛_구분} → {새_구분}  {str(이름)[:34]}")

# 드롭다운 교체
남길_검증 = [검증 for 검증 in 시트.data_validations.dataValidation
             if "G" not in str(검증.sqref)]
시트.data_validations.dataValidation = 남길_검증

새_검증 = DataValidation(type="list", formula1='"자동화 개발,보고서 생성"', allow_blank=True)
시트.add_data_validation(새_검증)
새_검증.add(f"G5:G{마지막_행}")

# 헤더 설명도 함께 고친다
시트.cell(4, 열["구분"]).comment = None
from openpyxl.comments import Comment
시트.cell(4, 열["구분"]).comment = Comment(
    "자동화 개발 / 보고서 생성 중에서 고르세요. "
    "홈페이지 '업무 자동화·솔루션 개발' 카드에서 이 기준으로 나뉘어 표시됩니다.",
    "로직플로우",
)

통합문서.save(엑셀_경로)
print(f"\n{바뀐_수}개 행의 구분을 바꿨습니다.")
